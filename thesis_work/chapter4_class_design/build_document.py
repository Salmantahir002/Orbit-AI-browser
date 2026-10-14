from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[2]
WORK = Path(__file__).resolve().parent

PARAGRAPHS = [
    "The class design of Orbit organises browser control, prompt-injection detection and AI assistance around TypeScript service classes. The React interface and several supporting operations use functions, while TypeScript interfaces define exchanged data. The accompanying diagram shows the principal classes and their relationships, identifying the security route module separately.",
    "BrowserRuntime maintains registered browsing targets and temporary element handles. Its registerTarget(), unregisterTarget() and invoke() operations manage guest targets and dispatch commands. CdpSession represents a guest-page connection through the Chrome DevTools Protocol. CdpSessionRegistry maintains one shared session per guest for BrowserRuntime and CdpInspectionService. The inspection service maintains capture state and provides watch(), forget() and capture() operations to collect webpage content for user-initiated scans.",
    "BrowserUIController manages menu surfaces, active downloads and zoom information. It owns a BrowserProfileStore containing history, bookmarks, download records and settings. The store provides recordVisit(), saveBookmark(), updateSettings() and snapshot(), while flush() saves the profile in a local JSON file. The React browser shell accesses these classes through the preload bridge.",
    "PromptPreprocessingService removes HTML where applicable and normalises case and whitespace. TextChunkingService divides text into identified, overlapping sections. The security route module coordinates these services and calls PromptClassifierService, which maintains model readiness and combines RuleBasedDetectorService with local Prompt Guard 2 inference. The rule detector returns suspicious patterns and matched evidence. Either detector can flag a chunk as malicious; if the model cannot load, classification continues with rules.",
    "FeatureExplanationService derives supporting evidence from text and matched phrases. The security route module combines this evidence with chunk results and the final decision for the analysis panels. These results are structured data contracts. The classifier also checks the chat prompt before LlmProviderManager forwards the request. This check covers the prompt field; optional page context, history and attachments are not automatically classified on that path.",
    "AgentTask coordinates Agent Mode and holds the task identifier, goal, target identifier, execution limit and progress. Each task creates one AgentWorkingMemory object to record actions, failures, retries, findings and milestones. Its snapshot() supplies bounded context to the planner. AgentTask also uses the shared AgentSessionMemory instance for recent terminated-task summaries during the renderer's lifetime.",
    "AgentPlannerService receives the goal, memory snapshot and page state through the local planning endpoint. Its buildMessages(), requestPlan() and parsePlan() operations prepare model input and validate the action queue through the tool registry. AgentTask applies approval checks and invokes BrowserRuntime through interprocess communication. The runtime executes commands and returns observations or verification results. This planning path does not automatically classify the goal or page observations with PromptClassifierService.",
    "ProviderGateway is the abstract class for AI-provider access. OpenAICompatibleGateway, AnthropicGateway and GeminiGateway implement listModels(), chatCompletion() and validateKey(). LlmProviderManager holds the active configuration and gateway, allowing chat and planning to share the provider contract. A factory selects the concrete gateway from the configuration. This is the main inheritance relationship; browser, detection and agent classes cooperate through ownership, shared references and service dependencies.",
]

doc = Document()
# Remove default template title rules and font-theme overrides before authoring.
for style in doc.styles:
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    for fonts in style.element.xpath('.//w:rFonts'):
        for theme in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme', 'csTheme'):
            fonts.attrib.pop(qn('w:'+theme), None)
section = doc.sections[0]
section.page_width = Mm(210)
section.page_height = Mm(297)
section.top_margin = section.bottom_margin = Mm(25.4)
section.left_margin = section.right_margin = Mm(25.4)
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(12)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
normal.element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.widow_control = True
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
title = doc.styles['Title']
title.font.name = 'Times New Roman'
title.font.size = Pt(14)
title.font.bold = True
title.font.color.rgb = RGBColor(0, 0, 0)
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(12)
title.paragraph_format.line_spacing = 1
title.paragraph_format.keep_with_next = True
heading = doc.add_paragraph('4.3 Class Design', style='Title')
heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = heading.runs[0]
run.font.name = 'Times New Roman'
run.font.size = Pt(14)
run.font.bold = True
run.font.color.rgb = RGBColor(0,0,0)
no_border = OxmlElement('w:pBdr')
for side in ('top', 'left', 'bottom', 'right', 'between'):
    item = OxmlElement('w:'+side)
    item.set(qn('w:val'), 'nil')
    no_border.append(item)
heading._p.get_or_add_pPr().append(no_border)
for text in PARAGRAPHS:
    doc.add_paragraph(text)
doc.core_properties.title = 'Orbit Chapter 4 Section 4.3 Class Design'
doc.core_properties.subject = 'Implemented class design of Orbit'
doc.core_properties.author = ''
doc.core_properties.keywords = 'Orbit, class design, system design'
doc.core_properties.comments = ''
output = ROOT / 'Chapter_4_Section_4_3_Class_Design.docx'
doc.save(output)
(WORK / 'section_4_3_class_design.md').write_text('# 4.3 Class Design\n\n' + '\n\n'.join(PARAGRAPHS) + '\n', encoding='utf-8')
print(f'Output: {output}')
print(f'Body words: {sum(len(p.split()) for p in PARAGRAPHS)}')
